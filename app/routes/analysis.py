import os

import fitz
import pytesseract
from docx import Document as DocxDocument
from flask import Blueprint, current_app, render_template
from flask_login import current_user, login_required
from PIL import Image, ImageOps

from app.models.document import Document
from app.services.ai.classifier import classify_document
from app.services.ai.recommendation import generate_recommendations
from app.services.ai.summarizer import (
    generate_summary,
    extract_key_information,
)
from app.services.ai.clause_detector import detect_clauses
from app.services.ai.missing_clause import detect_missing_clauses
from app.services.ai.risk_analysis import (
    analyze_risks,
    calculate_risk_score,
)

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

analysis_bp = Blueprint("analysis", __name__)


# ==========================================================
# OCR
# ==========================================================


def preprocess_image(image):
    """Prepare an image for OCR."""

    image = image.convert("RGB")

    image = image.resize(
        (image.width * 2, image.height * 2),
        Image.Resampling.LANCZOS,
    )

    image = ImageOps.grayscale(image)
    image = ImageOps.autocontrast(image)

    return image


def ocr_image(image):
    """Extract English text and confidence."""

    processed = preprocess_image(image)

    data = pytesseract.image_to_data(
        processed,
        lang="eng",
        config="--oem 3 --psm 3",
        output_type=pytesseract.Output.DICT,
    )

    text_parts = []
    confidences = []

    for i, text in enumerate(data["text"]):

        text = text.strip()

        if not text:
            continue

        try:
            conf = float(data["conf"][i])
        except (ValueError, TypeError):
            continue

        if conf >= 0:
            text_parts.append(text)
            confidences.append(conf)

    extracted = " ".join(text_parts)

    confidence = round(sum(confidences) / len(confidences), 1) if confidences else 0

    return extracted.strip(), confidence


# ==========================================================
# PDF
# ==========================================================


def extract_pdf_text(file_path):
    """Extract selectable PDF text."""

    text_parts = []

    with fitz.open(file_path) as pdf:

        for page in pdf:
            text_parts.append(page.get_text())

    return "\n".join(text_parts).strip()


def extract_scanned_pdf(file_path):
    """OCR scanned PDF."""

    text_parts = []
    confidence_values = []

    with fitz.open(file_path) as pdf:

        for page in pdf:

            pix = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False,
            )

            image = Image.frombytes(
                "RGB",
                [pix.width, pix.height],
                pix.samples,
            )

            text, conf = ocr_image(image)

            if text:
                text_parts.append(text)

            confidence_values.append(conf)

    confidence = (
        round(sum(confidence_values) / len(confidence_values), 1)
        if confidence_values
        else 0
    )

    return "\n\n".join(text_parts).strip(), confidence


# ==========================================================
# OTHER FILE TYPES
# ==========================================================


def extract_docx_text(file_path):

    document = DocxDocument(file_path)

    parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:

        for row in table.rows:

            row_text = " | ".join(cell.text.strip() for cell in row.cells)

            if row_text:
                parts.append(row_text)

    return "\n".join(parts).strip()


def extract_txt_text(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:

        return file.read().strip()


def extract_image_text(file_path):

    with Image.open(file_path) as image:

        return ocr_image(image)


# ==========================================================
# MAIN EXTRACTION
# ==========================================================


def analyze_file(file_path):

    extension = file_path.rsplit(".", 1)[1].lower()

    result = {
        "text": "",
        "method": "Unknown",
        "pages": 1,
        "ocr_used": False,
        "ocr_pages": 0,
        "confidence": None,
    }

    # ---------------- PDF ----------------

    if extension == "pdf":

        with fitz.open(file_path) as pdf:
            result["pages"] = len(pdf)

        text = extract_pdf_text(file_path)

        if len(text.strip()) >= 50:

            result["text"] = text
            result["method"] = "PDF Text Extraction"

        else:

            text, confidence = extract_scanned_pdf(file_path)

            result["text"] = text
            result["method"] = "Tesseract OCR"
            result["ocr_used"] = True
            result["ocr_pages"] = result["pages"]
            result["confidence"] = confidence

    # ---------------- DOCX ----------------

    elif extension == "docx":

        result["text"] = extract_docx_text(file_path)
        result["method"] = "DOCX Text Extraction"

    # ---------------- TXT ----------------

    elif extension == "txt":

        result["text"] = extract_txt_text(file_path)
        result["method"] = "TXT Text Extraction"

    # ---------------- IMAGE ----------------

    elif extension in {"png", "jpg", "jpeg"}:

        text, confidence = extract_image_text(file_path)

        result["text"] = text
        result["method"] = "Tesseract OCR"
        result["ocr_used"] = True
        result["ocr_pages"] = 1
        result["confidence"] = confidence

    return result


# ==========================================================
# ANALYSIS HOME
# ==========================================================


@analysis_bp.route("/analysis")
@login_required
def analysis_home():

    documents = (
        Document.query.filter_by(user_id=current_user.id)
        .order_by(Document.upload_date.desc())
        .all()
    )

    return render_template(
        "analysis/analysis_home.html",
        documents=documents,
    )


# ==========================================================
# ANALYZE DOCUMENT
# ==========================================================


@analysis_bp.route("/analysis/<int:document_id>")
@login_required
def analyze(document_id):

    document = Document.query.filter_by(
        id=document_id,
        user_id=current_user.id,
    ).first_or_404()

    upload_folder = os.path.join(
        current_app.root_path,
        "static",
        "uploads",
    )

    file_path = os.path.join(
        upload_folder,
        document.filename,
    )

    result = analyze_file(file_path)

    extracted_text = result.get("text", "")

    if not extracted_text:

        extracted_text = "No readable text could be extracted from this document."

    # ======================================================
    # AI MODULES
    # ======================================================

    classification = classify_document(extracted_text)

    document_type = classification["type"]
    classification_confidence = classification["confidence"]
    classification_description = classification["description"]

    summary = generate_summary(extracted_text)

    key_information = extract_key_information(extracted_text)

    clauses = detect_clauses(extracted_text)

    # NEW FEATURE
    missing_clauses = detect_missing_clauses(extracted_text)

    legal_documents = {
        "Contract / Agreement",
        "Will / Testament",
        "Lease / Rental Agreement",
        "Employment Agreement",
        "Non-Disclosure Agreement",
    }

    if document_type in legal_documents:

        risks = analyze_risks(extracted_text)
        risk_score = calculate_risk_score(risks)

    else:

        risks = []
        risk_score = None

    if risk_score is None:
        risk_level = "Not Applicable"

    elif risk_score >= 60:
        risk_level = "High"

    elif risk_score >= 30:
        risk_level = "Medium"

    else:
        risk_level = "Low"

    recommendations = generate_recommendations(
        document_type,
        risk_score,
        missing_clauses,
        risks,
    )

    # ======================================================
    # Render
    # ======================================================

    return render_template(
        "analysis/analysis.html",
        document=document,
        # Extraction
        extracted_text=extracted_text,
        extraction_method=result["method"],
        pages=result["pages"],
        ocr_used=result["ocr_used"],
        ocr_pages=result["ocr_pages"],
        confidence=result["confidence"],
        # AI Results
        classification=classification,
        document_type=document_type,
        classification_confidence=classification_confidence,
        classification_description=classification_description,
        summary=summary,
        key_information=key_information,
        clauses=clauses,
        missing_clauses=missing_clauses,
        risks=risks,
        risk_score=risk_score,
        risk_level=risk_level,
        recommendations=recommendations,
    )
